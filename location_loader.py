"""Read location rows from the user workbook without modifying it."""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook


class LocationDataError(ValueError):
    pass


@dataclass(frozen=True)
class Location:
    province: str
    city: str
    county: str

    @property
    def full_name(self) -> str:
        return f"{self.province}{self.city}{self.county}"

    @property
    def spoken_name(self) -> str:
        return f"{self.city}{self.county}"


def load_locations(path: str | Path) -> list[Location]:
    workbook_path = Path(path)
    if not workbook_path.is_file():
        raise FileNotFoundError(f"地点 Excel 不存在：{workbook_path}")
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)
        try:
            raw_headers = next(rows)
        except StopIteration as exc:
            raise LocationDataError("地点 Excel 是空文件") from exc
        headers = {str(value).strip(): index for index, value in enumerate(raw_headers) if value is not None}
        required = ("省份", "州市", "县区")
        missing = [name for name in required if name not in headers]
        if missing:
            raise LocationDataError(f"地点 Excel 缺少列：{'、'.join(missing)}")
        locations = []
        for excel_row, values in enumerate(rows, start=2):
            selected = [values[headers[name]] if headers[name] < len(values) else None for name in required]
            if all(value is None or not str(value).strip() for value in selected):
                continue
            if any(value is None or not str(value).strip() for value in selected):
                raise LocationDataError(f"地点 Excel 第 {excel_row} 行存在空白地点字段")
            locations.append(Location(*(str(value).strip() for value in selected)))
        if not locations:
            raise LocationDataError("地点 Excel 没有有效数据行")
        return locations
    finally:
        workbook.close()


def load_location(path: str | Path, row_number: int = 1) -> Location:
    if row_number < 1:
        raise LocationDataError("location-row 必须从 1 开始")
    locations = load_locations(path)
    if row_number > len(locations):
        raise LocationDataError(f"location-row={row_number} 超出范围，当前只有 {len(locations)} 条地点")
    return locations[row_number - 1]
